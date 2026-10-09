"""Convert an Obsidian markdown draft (PRD or Decision Doc) to a Word document."""
import re
import sys
from pathlib import Path
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS_DIR = Path(__file__).parent.parent / "docs" / "output"


def _add_table(doc, lines: list[str]):
    rows = [re.split(r"\s*\|\s*", l.strip().strip("|")) for l in lines if not re.match(r"^\|[-| ]+\|$", l)]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, cell in enumerate(row):
            p = table.rows[r_idx].cells[c_idx].paragraphs[0]
            run = p.add_run(cell.strip())
            if r_idx == 0:
                run.bold = True


def convert(md_path: str) -> Path:
    path = Path(md_path)
    doc = docx.Document()

    # Remove YAML frontmatter
    text = path.read_text(encoding="utf-8")
    text = re.sub(r"^---.*?---\n", "", text, flags=re.DOTALL)

    lines = text.splitlines()
    table_buffer = []
    i = 0

    while i < len(lines):
        line = lines[i]

        # Flush table buffer when a non-table line appears
        if table_buffer and not line.startswith("|"):
            _add_table(doc, table_buffer)
            table_buffer = []

        if line.startswith("|"):
            table_buffer.append(line)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- [ ] ") or line.startswith("- [x] "):
            doc.add_paragraph(line[6:], style="List Bullet")
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip() == "---":
            doc.add_paragraph()
        elif line.strip():
            doc.add_paragraph(line)

        i += 1

    if table_buffer:
        _add_table(doc, table_buffer)

    out_path = DOCS_DIR / (path.stem + ".docx")
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Saved: {out_path}")
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python export_word.py <path_to_md>", file=sys.stderr)
        sys.exit(1)
    convert(sys.argv[1])
